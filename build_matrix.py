"""Build the Responsibility Matrix HTML page from the Excel workbook.

Usage:  python build_matrix.py <workbook.xlsx> [output.html]
Re-run whenever the Excel file is updated.
"""
import base64, html, re, sys, zipfile
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

RED = "FFFF0000"


def is_red(font):
    try:
        return font is not None and font.color is not None and font.color.type == "rgb" and font.color.rgb == RED
    except Exception:
        return False


def span(text, bold=False, red=False, strike=False):
    t = html.escape(text).replace("\n", "<br>")
    if not t.strip():
        return t
    if strike:
        t = f"<s>{t}</s>"
    if red:
        t = f'<span class="red">{t}</span>'
    if bold:
        t = f"<b>{t}</b>"
    return t


def cell_html(cell):
    v = cell.value
    if v is None:
        return ""
    base = cell.font
    if isinstance(v, CellRichText):
        out = []
        for part in v:
            if isinstance(part, TextBlock):
                f = part.font
                out.append(span(part.text, bool(f.b), is_red(f) or (f.color is None and is_red(base)), bool(f.strike)))
            else:
                out.append(span(str(part), bool(base.b), is_red(base), bool(base.strike)))
        return "".join(out).strip()
    return span(str(v).strip(), False, is_red(base), bool(base.strike))


def main(src, dst):
    wb = openpyxl.load_workbook(src, rich_text=True)
    ws = wb.active

    # logo embedded in the workbook
    logo = ""
    with zipfile.ZipFile(src) as z:
        imgs = [n for n in z.namelist() if n.startswith("xl/media/")]
        if imgs:
            ext = imgs[0].rsplit(".", 1)[-1].lower()
            logo = f"data:image/{'jpeg' if ext in ('jpg','jpeg') else ext};base64," + base64.b64encode(z.read(imgs[0])).decode()

    updated = cell_html(ws["C5"])
    title = cell_html(ws["B6"])
    store = cell_html(ws["B7"])
    plain_title = re.sub(r"<[^>]+>", " ", title)

    rows = []
    for r in range(11, ws.max_row + 1):
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        if b is None and c is None:
            continue
        code = str(b).strip() if b is not None else ""
        # division header rows: 2-digit code (e.g. "00", 10, 21)
        if re.fullmatch(r"\d{1,2}", code):
            rows.append(f'<tr class="div"><td class="c">{int(code):02d}</td><td colspan="8">{html.escape(str(c).strip())}</td></tr>')
            continue
        flags = {}
        tds = [f'<td class="c nw">{cell_html(ws.cell(r, 2))}</td>', f"<td>{cell_html(ws.cell(r, 3))}</td>"]
        for col, cls in ((4, "gc"), (5, "gc"), (6, "ow"), (7, "ow")):
            val = str(ws.cell(r, col).value or "").strip().upper()
            flags["defg"[col - 4]] = "X" if val == "X" else ""
            tds.append(f'<td class="c x {cls}">X</td>' if val == "X" else '<td class="c x"></td>')
        tds.append(f'<td class="c">{cell_html(ws.cell(r, 8))}</td>')
        tds.append(f'<td class="notes">{cell_html(ws.cell(r, 9))}</td>')
        sub = str(ws.cell(r, 10).value or "").strip().upper()
        tds.append('<td class="c x sub">Y</td>' if sub == "Y" else '<td class="c x"></td>')
        vendor = " ".join(str(ws.cell(r, 8).value or "").split())
        attrs = "".join(f' data-{k}="{v}"' for k, v in flags.items()) + f' data-h="{html.escape(vendor)}"'
        rows.append(f'<tr class="item"{attrs}>' + "".join(tds) + "</tr>")
    # blank trailing lines, like the printed sheet
    rows += ['<tr class="blank">' + "<td></td>" * 9 + "</tr>"] * 5

    page = TEMPLATE.format(
        page_title=html.escape(" ".join(plain_title.split())),
        logo=f'<img src="{logo}" alt="McCoy\'s Building Supply logo">' if logo else "",
        updated=updated, title=title, store=store, rows="\n".join(rows),
    )
    with open(dst, "w", encoding="utf-8") as fh:
        fh.write(page)
    print("wrote", dst, len(rows), "rows")


TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{page_title}</title>
<style>
  :root {{
    --green: #548235; --sect: #E2EFDA; --gc: #FFE699; --owner: #A9D08E;
    --sub: #B4C6E7; --yellow: #FFFF00; --line: #000;
  }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; background: #fff; color: #000; font-family: "Times New Roman", Times, serif; font-size: 13px; }}
  .page {{ max-width: 1400px; margin: 0 auto; padding: 24px 16px 48px; }}
  .top {{ display: flex; justify-content: space-between; align-items: flex-end; }}
  .top img {{ height: 78px; display: block; }}
  .corp {{ font-family: Arial, Helvetica, sans-serif; font-size: 11px; text-align: center; line-height: 1.35; }}
  @media (max-width: 1150px) {{ .scroll {{ overflow-x: auto; }} }}
  table {{ border-collapse: collapse; width: 100%; min-width: 1100px; border: 2px solid var(--line); }}
  td, th {{ border: 1px solid var(--line); padding: 2px 4px; vertical-align: middle; font-weight: normal; }}
  .updated {{ text-align: center; font-weight: bold; padding: 4px 0 3px; }}
  .red {{ color: #FF0000; }}
  .updated .red {{ color: #FF0000; }}
  .banner th {{ background: var(--green); color: var(--yellow); font-weight: bold; font-size: 17px; padding: 6px; border-bottom: 2px solid var(--line); }}
  .store td {{ text-align: center; padding: 5px; line-height: 1.4; }}
  thead tr.h th {{ background: var(--green); color: var(--yellow); text-align: center; position: sticky; z-index: 1; }}
  thead tr.h1 th {{ top: 0; }}
  thead tr.h2 th {{ top: 22px; }}
  thead tr.h th.gch {{ background: var(--gc); color: #000; }}
  thead tr.h th.owh {{ background: var(--owner); color: #000; }}
  thead tr.h th.subh {{ background: var(--sub); color: #000; }}
  .c {{ text-align: center; }}
  .nw {{ white-space: nowrap; }}
  .x {{ width: 58px; }}
  td.gc {{ background: var(--gc); }}
  td.ow {{ background: var(--owner); }}
  td.sub {{ background: var(--sub); }}
  tr.div td {{ background: var(--sect); font-weight: bold; }}
  tr.blank td {{ height: 18px; }}
  tr.filters th {{ background: #f3f7ef; padding: 3px 2px; position: sticky; z-index: 1; }}
  tr.filters select {{ width: 100%; font: inherit; font-size: 12px; padding: 1px; border: 1px solid #8aa66f; border-radius: 2px; background: #fff; }}
  tr.filters select.on {{ background: #FFF2CC; border-color: #548235; font-weight: bold; }}
  tr.filters .info {{ font-size: 12px; text-align: left; white-space: nowrap; }}
  tr.filters button {{ font: inherit; font-size: 12px; margin-left: 8px; padding: 1px 8px; border: 1px solid #548235; background: #fff; color: #385723; border-radius: 3px; cursor: pointer; }}
  tr.filters button:disabled {{ opacity: .4; cursor: default; }}
  tr.hide {{ display: none; }}
  .ms {{ position: relative; }}
  .ms-btn {{ width: 100%; margin: 0 !important; text-align: left; font-weight: normal; color: #000 !important; border: 1px solid #8aa66f !important; border-radius: 2px !important; padding: 1px 16px 1px 4px !important; background: #fff !important; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; position: relative; }}
  .ms-btn::after {{ content: "▾"; position: absolute; right: 4px; top: 0; }}
  .ms-btn.on {{ background: #FFF2CC !important; border-color: #548235 !important; font-weight: bold; }}
  .ms-panel {{ position: absolute; top: calc(100% + 2px); left: 0; z-index: 5; width: 250px; background: #fff; border: 1px solid #548235; border-radius: 3px; box-shadow: 0 4px 14px rgba(0,0,0,.18); text-align: left; font-weight: normal; }}
  .ms-actions {{ display: flex; justify-content: space-between; padding: 5px 6px; border-bottom: 1px solid #d6e4c8; background: #f3f7ef; }}
  .ms-actions button {{ margin: 0 !important; }}
  .ms-list {{ max-height: 320px; overflow-y: auto; padding: 3px 0; }}
  .ms-list label {{ display: flex; align-items: center; gap: 6px; padding: 3px 8px; font-size: 12px; cursor: pointer; white-space: normal; }}
  .ms-list label:hover {{ background: #E2EFDA; }}
  .ms-list label.sep {{ border-bottom: 1px solid #e3e3e3; font-style: italic; }}
  .ms-list input {{ margin: 0; accent-color: #548235; }}
  tbody tr:not(.div):not(.blank):hover td {{ outline: 1px solid #548235; outline-offset: -1px; }}
  .notes {{ line-height: 1.35; }}
  .foot {{ text-align: center; font-size: 11px; margin-top: 16px; }}
  @media print {{
    .page {{ max-width: none; padding: 0; }}
    .scroll {{ overflow: visible; }}
    table {{ min-width: 0; font-size: 9px; }}
    thead tr.h th {{ position: static; }}
    tr {{ break-inside: avoid; }}
    tr.filters {{ display: none; }}
    * {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
  }}
</style>
</head>
<body>
<div class="page">
  <div class="top">
    <div>{logo}</div>
    <div class="corp">McCoy Corporation<br><b>Store Development Dept.</b><br>1350 IH 35 North<br>San Marcos, Texas 78667<br>(512) 395-6644</div>
  </div>
  <div class="updated">{updated}</div>
  <div class="scroll">
  <table>
    <colgroup>
      <col style="width:66px"><col style="width:24%"><col><col><col><col>
      <col style="width:145px"><col style="width:46%"><col style="width:66px">
    </colgroup>
    <thead>
      <tr class="banner"><th colspan="9">{title}</th></tr>
      <tr class="store"><td colspan="9">{store}</td></tr>
      <tr class="h h1">
        <th rowspan="2">CSI<br>Division</th>
        <th rowspan="2">Description</th>
        <th colspan="2" class="gch">GC</th>
        <th colspan="2" class="owh">Owner</th>
        <th rowspan="2">Use Specific<br>Vendor/Supplier</th>
        <th rowspan="2">Comments &amp; Notes</th>
        <th class="subh">Submittal</th>
      </tr>
      <tr class="h h2">
        <th class="gch">Furnish</th><th class="gch">Install</th>
        <th class="owh">Furnish</th><th class="owh">Install</th>
        <th class="subh">Required</th>
      </tr>
      <tr class="filters">
        <th></th>
        <th class="info" style="text-align:right">Filter:</th>
        <th><select data-col="d" aria-label="Filter GC Furnish"><option value="">All</option><option value="X">X</option><option value="-">Blank</option></select></th>
        <th><select data-col="e" aria-label="Filter GC Install"><option value="">All</option><option value="X">X</option><option value="-">Blank</option></select></th>
        <th><select data-col="f" aria-label="Filter Owner Furnish"><option value="">All</option><option value="X">X</option><option value="-">Blank</option></select></th>
        <th><select data-col="g" aria-label="Filter Owner Install"><option value="">All</option><option value="X">X</option><option value="-">Blank</option></select></th>
        <th><div class="ms" id="vendor"><button type="button" class="ms-btn" aria-haspopup="true" aria-expanded="false" aria-label="Filter Vendor/Supplier">All</button><div class="ms-panel" hidden><div class="ms-actions"><button type="button" data-act="all">Select all</button><button type="button" data-act="none">Clear</button></div><div class="ms-list"></div></div></div></th>
        <th class="info"><span id="count"></span><button type="button" id="clear" disabled>Clear filters</button></th>
        <th></th>
      </tr>
    </thead>
    <tbody>
{rows}
    </tbody>
  </table>
  </div>
</div>
<script>
(function () {{
  var items = Array.prototype.slice.call(document.querySelectorAll("tbody tr.item"));
  var rows = Array.prototype.slice.call(document.querySelectorAll("tbody tr"));
  var selects = Array.prototype.slice.call(document.querySelectorAll("tr.filters select"));
  var count = document.getElementById("count"), clear = document.getElementById("clear");

  // vendor list, alphabetical
  var vendors = {{}};
  items.forEach(function (r) {{ if (r.dataset.h) vendors[r.dataset.h] = 1; }});
  var ms = document.getElementById("vendor"), msBtn = ms.querySelector(".ms-btn"),
      msPanel = ms.querySelector(".ms-panel"), msList = ms.querySelector(".ms-list");
  function addBox(value, text, cls) {{
    var l = document.createElement("label"); if (cls) l.className = cls;
    var cb = document.createElement("input"); cb.type = "checkbox"; cb.value = value;
    l.appendChild(cb); l.appendChild(document.createTextNode(text)); msList.appendChild(l);
  }}
  addBox("-", "(Blank – no vendor listed)", "sep");
  Object.keys(vendors).sort(function (a, b) {{ return a.localeCompare(b); }}).forEach(function (v) {{ addBox(v, v); }});
  var boxes = Array.prototype.slice.call(msList.querySelectorAll("input"));
  function picked() {{ return boxes.filter(function (b) {{ return b.checked; }}).map(function (b) {{ return b.value; }}); }}
  function openPanel(open) {{ msPanel.hidden = !open; msBtn.setAttribute("aria-expanded", open); }}
  msBtn.addEventListener("click", function (e) {{ e.stopPropagation(); openPanel(msPanel.hidden); }});
  msPanel.addEventListener("click", function (e) {{ e.stopPropagation(); }});
  document.addEventListener("click", function () {{ openPanel(false); }});
  document.addEventListener("keydown", function (e) {{ if (e.key === "Escape") openPanel(false); }});
  msList.addEventListener("change", apply);
  ms.querySelector("[data-act=all]").addEventListener("click", function () {{ boxes.forEach(function (b) {{ b.checked = true; }}); apply(); }});
  ms.querySelector("[data-act=none]").addEventListener("click", function () {{ boxes.forEach(function (b) {{ b.checked = false; }}); apply(); }});

  function apply() {{
    var active = selects.filter(function (s) {{ return s.value; }});
    selects.forEach(function (s) {{ s.classList.toggle("on", !!s.value); }});
    // vendor checklist: nothing ticked (or everything ticked) = no vendor filter
    var vp = picked(), vOn = vp.length > 0 && vp.length < boxes.length;
    msBtn.classList.toggle("on", vOn);
    msBtn.textContent = !vOn ? "All" : vp.length === 1 ? (vp[0] === "-" ? "(Blank)" : vp[0]) : vp.length + " selected";
    msBtn.title = vOn ? vp.map(function (v) {{ return v === "-" ? "(Blank)" : v; }}).join(", ") : "";
    var nActive = active.length + (vOn ? 1 : 0);
    var shown = 0;
    items.forEach(function (r) {{
      var ok = active.every(function (s) {{
        var v = r.dataset[s.dataset.col] || "";
        return s.value === "-" ? v === "" : v === s.value;
      }}) && (!vOn || vp.indexOf(r.dataset.h || "-") !== -1);
      r.classList.toggle("hide", !ok);
      if (ok) shown++;
    }});
    // hide division headings with nothing showing under them, and the blank lines while filtering
    var div = null, any = false;
    rows.forEach(function (r) {{
      if (r.classList.contains("div")) {{ if (div) div.classList.toggle("hide", !any); div = r; any = false; }}
      else if (r.classList.contains("item")) {{ if (!r.classList.contains("hide")) any = true; }}
      else if (r.classList.contains("blank")) {{ r.classList.toggle("hide", nActive > 0); }}
    }});
    if (div) div.classList.toggle("hide", !any);
    count.textContent = nActive ? "Showing " + shown + " of " + items.length + " items" : items.length + " items";
    clear.disabled = !nActive;
  }}
  selects.forEach(function (s) {{ s.addEventListener("change", apply); }});
  clear.addEventListener("click", function () {{ selects.forEach(function (s) {{ s.value = ""; }}); boxes.forEach(function (b) {{ b.checked = false; }}); apply(); }});

  // keep the header rows stacked correctly while scrolling
  function stick() {{
    var h1 = document.querySelector("tr.h1"), h2 = document.querySelector("tr.h2");
    h1.querySelectorAll("th").forEach(function (th) {{ th.style.top = "0px"; }});
    h2.querySelectorAll("th").forEach(function (th) {{ th.style.top = (h1.offsetHeight) + "px"; }});
    document.querySelectorAll("tr.filters th").forEach(function (th) {{ th.style.top = (h1.offsetHeight + h2.offsetHeight) + "px"; }});
  }}
  window.addEventListener("resize", stick);
  stick();
  apply();
}})();
</script>
</body>
</html>
"""

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "index.html")
